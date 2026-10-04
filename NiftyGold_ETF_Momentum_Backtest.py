#!/usr/bin/env python3
# =============================================================================
# NIFTY:GOLD RATIO + ETF MOMENTUM ROTATION — COMBINED BACKTEST
# Single self-contained script. Paste into a GitHub repo and run:
#     python NiftyGold_ETF_Momentum_Backtest.py
# It backtests the same rules as the Tradetron strategy "Nifty Gold and ETF Momentum 1L":
#     Part A (49%)  NIFTYBEES + GOLDBEES, weights set by the ratio NIFTYBEES / GOLDBEES.
#                   Gold weight = (ratio - 1.25) / (6.25 - 1.25), limited to 0-100%.
#                   Buys on day 1; reshuffles only when the ratio has moved 1.0 from the
#                   ratio at the last purchase / reshuffle. Only the difference is traded.
#     Part B (49%)  ETF momentum rotation on 50 ETFs. Monthly (first trading day):
#                   Nifty 50 <= its 200-DMA -> sell all, stay in cash. Otherwise sell
#                   holdings ranked below 10, buy top-6 names that have 3M > 0 and are
#                   above their own 200-DMA, 1/6 of Part B value each, whole units.
#                   Hard stop -15% from entry and trailing stop -20% from peak, checked daily.
#     Cash (2%)     kept idle.
# Each part compounds on its own (no rebalancing between the parts).
# Outputs one Excel workbook + one self-contained dark HTML report.
# =============================================================================

import subprocess, sys

def _ensure(pkg, import_as=None):
    try:
        __import__(import_as or pkg)
    except ImportError:
        print(f"Installing {pkg}...")
        subprocess.check_call([sys.executable, "-m", "pip", "install", pkg, "--quiet"])

for _p, _i in [("yfinance", "yfinance"), ("pandas", "pandas"), ("numpy", "numpy"),
               ("openpyxl", "openpyxl"), ("plotly", "plotly")]:
    _ensure(_p, _i)

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

import warnings
warnings.filterwarnings("ignore")

import os, math, webbrowser
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
import plotly.graph_objects as go
import plotly.offline as pyo

# =============================================================================
# USER INPUTS  (on GitHub these come from the "Run workflow" form)
# =============================================================================
START_DATE = os.getenv("BT_START", "").strip() or "2015-01-01"
END_DATE   = os.getenv("BT_END", "").strip()                    # "" = latest available
CAPITAL    = float(os.getenv("BT_CAPITAL", "").strip() or 100000)
STOP_BASIS = (os.getenv("BT_STOP_BASIS", "").strip() or "close").lower()   # "close" or "low"
COST_PCT   = float(os.getenv("BT_COST_PCT", "").strip() or 0.10)  # % of trade value, each side
                                                                 # (charges + slippage together)
DP_CHARGE      = 15.34    # Rs per sell transaction (depository charge); 0 to ignore
CASH_YIELD_PA  = 0.0      # idle cash earns nothing in a broker account; e.g. 0.05 for a liquid fund
RF_PA          = 0.06     # risk-free rate used in the Sharpe ratio

# Split of the capital (matches the Tradetron strategy: 49k + 49k + 2k per 1 lakh)
NG_SHARE, MOM_SHARE, CASH_SHARE = 0.49, 0.49, 0.02

# ---- Part A: Nifty:Gold ratio allocation ----
NG_MIN, NG_MAX, NG_STEP = 1.25, 6.25, 1.0
NG_NIFTY, NG_GOLD = "NIFTYBEES", "GOLDBEES"

# ---- Part B: ETF momentum rotation (same numbers as the live tracker) ----
N_HOLD      = 6
HOLD_RANK   = 10
LB          = [21, 63, 126, 252]
WEIGHTS     = [0.15, 0.40, 0.30, 0.15]
DMA_PERIOD  = 200
MIN_HISTORY = 262
SL_PCT      = 0.15
TRAIL_PCT   = 0.20
STALE_DAYS  = 10          # an ETF with no price for this many days is not ranked / bought
REGIME_INDEX, REGIME_LABEL, REGIME_DMA = "^NSEI", "Nifty 50", 200
BENCH, BENCH_LABEL = "^CRSLDX", "Nifty 500"
JUMP_FLAG   = 0.35        # a one-day move bigger than this cannot be real for an ETF (20% price band)
GLITCH_MAX_DAYS = 10      # a wrong-price patch that returns to normal within this many days is repaired
SPLIT_FACTORS   = [2, 2.5, 4, 5, 10, 20, 25, 50, 100, 1000]

UNIVERSE = [
    ("NIFTYBEES", "Broad_Equity"), ("JUNIORBEES", "Broad_Equity"), ("MID150BEES", "Broad_Equity"),
    ("MIDSELIETF", "Broad_Equity"), ("MONIFTY500", "Broad_Equity"), ("HDFCSML250", "Broad_Equity"),
    ("MOM100", "Broad_Equity"),
    ("BANKBEES", "Banking_Finance"), ("PSUBNKBEES", "Banking_Finance"), ("PVTBANIETF", "Banking_Finance"),
    ("FINIETF", "Banking_Finance"),
    ("ITBEES", "Sector"), ("PHARMABEES", "Sector"), ("AUTOBEES", "Sector"), ("CONSUMBEES", "Sector"),
    ("CONSUMER", "Sector"), ("FMCGIETF", "Sector"), ("CHEMICAL", "Sector"), ("METALIETF", "Sector"),
    ("OILIETF", "Sector"), ("COMMOIETF", "Sector"), ("INFRAIETF", "Sector"), ("CPSEETF", "Sector"),
    ("ICICIB22", "Sector"), ("MOCAPITAL", "Sector"), ("MODEFENCE", "Sector"), ("MOREALTY", "Sector"),
    ("MOTOUR", "Sector"), ("GROWWPOWER", "Sector"), ("GROWWRAIL", "Sector"), ("GROWWHOSPI", "Sector"),
    ("DIVOPPBEES", "Thematic"), ("EVINDIA", "Thematic"), ("INTERNET", "Thematic"), ("MNC", "Thematic"),
    ("SELECTIPO", "Thematic"), ("TOP10ADD", "Thematic"),
    ("LOWVOLIETF", "Factor"), ("NV20IETF", "Factor"), ("QUAL30IETF", "Factor"), ("NIFTYQLITY", "Factor"),
    ("HDFCGROWTH", "Factor"), ("MOM50", "Factor"), ("MOMENTUM50", "Factor"),
    ("MOVALUE", "Factor"),
    ("GOLDBEES", "Gold_Silver"), ("SILVERBEES", "Gold_Silver"),
    ("LTGILTBEES", "Bonds"), ("EBBETF0430", "Bonds"), ("GILT5YBEES", "Bonds"),
]
CATEGORY = {s: c for s, c in UNIVERSE}

OUTPUT_DIR = Path(__file__).parent / "Combined_Backtest_Output"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
EXCEL_PATH = OUTPUT_DIR / "NiftyGold_ETF_Momentum_Backtest.xlsx"
HTML_PATH  = OUTPUT_DIR / "NiftyGold_ETF_Momentum_Backtest.html"


# =============================================================================
# DATA
# =============================================================================
def _clean(df):
    df = df[["Open", "High", "Low", "Close"]].copy()
    df.index = pd.to_datetime(df.index).tz_localize(None)
    df = df[~df.index.duplicated(keep="last")].sort_index()
    df = df[df["Close"].notna() & (df["Close"] > 0)]
    for c in ["Open", "High", "Low"]:
        df.loc[~(df[c] > 0), c] = np.nan
    return df


def load_data():
    """Returns (data: {symbol: OHLC DataFrame}, regime close Series, benchmark close Series, missing list)."""
    if os.getenv("BT_MOCK") == "1":
        return _mock_data()
    import yfinance as yf
    syms = [s for s, _ in UNIVERSE]
    ytk = [s + ".NS" for s in syms] + [BENCH, REGIME_INDEX]
    print(f"Downloading daily history for {len(syms)} ETFs + {REGIME_LABEL} + {BENCH_LABEL}...")
    raw = yf.download(ytk, period="max", interval="1d", auto_adjust=True,
                      group_by="ticker", threads=True, progress=False)

    def frame_of(tkr):
        try:
            if isinstance(raw.columns, pd.MultiIndex) and tkr in raw.columns.get_level_values(0):
                df = _clean(raw[tkr])
                return df if len(df) else None
        except Exception:
            return None
        return None

    data, missing = {}, []
    for s in syms:
        df = frame_of(s + ".NS")
        if df is None:
            missing.append(s)
        else:
            data[s] = df
    reg = frame_of(REGIME_INDEX)
    ben = frame_of(BENCH)
    print(f"{len(data)} ETFs loaded, {len(missing)} with no data"
          + (f" ({', '.join(missing)})" if missing else ""))
    return data, (reg["Close"] if reg is not None else None), (ben["Close"] if ben is not None else None), missing


def _mock_data():
    """Synthetic prices for an offline dry-run (BT_MOCK=1). Not real results."""
    rng = np.random.default_rng(7)
    idx = pd.bdate_range(end=datetime.now().date(), periods=4200)
    mkt = rng.normal(0.0004, 0.010, len(idx))
    # a bear phase so the regime switch and the stops get exercised
    mkt[1500:1700] -= 0.0030
    mkt[3000:3120] -= 0.0035
    data = {}
    for i, (s, _) in enumerate(UNIVERSE):
        start = 0 if i < 12 else int(rng.integers(300, 3700))
        n = len(idx) - start
        beta = 0.6 + (i % 7) * 0.12
        r = beta * mkt[start:] + rng.normal(0.00005 * (i % 5), 0.009, n)
        if s == "GOLDBEES":
            r = rng.normal(0.00045, 0.008, n) - 0.3 * mkt[start:]
        close = (30 if s == "GOLDBEES" else 100) * np.exp(np.cumsum(r))
        spread = np.abs(rng.normal(0.006, 0.004, n))
        data[s] = pd.DataFrame({"Open": close * (1 + rng.normal(0, 0.003, n)),
                                "High": close * (1 + spread), "Low": close * (1 - spread),
                                "Close": close}, index=idx[start:])
    if os.getenv("BT_MOCK_GLITCH") == "1":      # reproduce the Dec-2019 Yahoo errors + one missed split
        for s_, k in (("NIFTYBEES", 10), ("GOLDBEES", 100), ("BANKBEES", 10)):
            data[s_].iloc[-1700:-1698] = data[s_].iloc[-1700:-1698].values / k
        data["ITBEES"].iloc[:-900] = data["ITBEES"].iloc[:-900].values * 10
    reg = pd.Series(8000 * np.exp(np.cumsum(mkt)), index=idx)
    ben = pd.Series(7000 * np.exp(np.cumsum(mkt + rng.normal(0, 0.002, len(idx)))), index=idx)
    print(f"[MOCK] generated {len(data)} ETFs + regime index + benchmark (synthetic, not real results)")
    return data, reg, ben, []


def _snap(f):
    """Snap a price ratio to the nearest usual split factor (or its inverse) when it is within ~12%."""
    inv = f < 1
    g = 1 / f if inv else f
    best = min(SPLIT_FACTORS, key=lambda k: abs(math.log(g / k)))
    if abs(math.log(g / best)) < 0.12:
        g = float(best)
    return 1 / g if inv else g


def repair_prices(df):
    """Fix Yahoo split errors in one ETF's history. Returns (repaired df, notes, needs_review).
    Two cases, both recognised by a one-day move no real ETF can make (beyond JUMP_FLAG):
      1. A short patch of wrong prices that snaps back (e.g. two days shown at 1/10th or 1/100th
         around a unit split): that patch is scaled back onto the surrounding level.
      2. A jump that never comes back (a split Yahoo did not adjust for): all earlier prices are
         rescaled so the series is continuous. This case is marked for a manual look."""
    df = df.copy()
    cols = [df.columns.get_loc(c) for c in ["Open", "High", "Low", "Close"]]
    c = df["Close"].values.astype(float).copy()
    hi, lo = 1 + JUMP_FLAG, 1 / (1 + JUMP_FLAG)
    notes, review, t, guard = [], False, 1, 0
    while t < len(c) and guard < 40:
        r = c[t] / c[t - 1]
        if lo <= r <= hi:
            t += 1
            continue
        guard += 1
        base, t2 = c[t - 1], None
        for j in range(t + 1, min(t + 1 + GLITCH_MAX_DAYS, len(c))):
            if lo <= c[j] / base <= hi:
                t2 = j
                break
        if t2 is not None:
            f = _snap(base / float(np.median(c[t:t2])))
            df.iloc[t:t2, cols] = df.iloc[t:t2, cols].values * f
            c[t:t2] *= f
            notes.append(f"{df.index[t].date()} to {df.index[t2 - 1].date()}: {t2 - t} day(s) of wrong prices multiplied by {f:g}")
            t = t2
        else:
            f = _snap(c[t] / base)
            df.iloc[:t, cols] = df.iloc[:t, cols].values * f
            c[:t] *= f
            notes.append(f"{df.index[t].date()}: jump that never reverted, all earlier prices multiplied by {f:g}")
            review = True
            t += 1
    return df, notes, review


def repair_all(data):
    out, repairs = {}, {}
    for s, df in data.items():
        out[s], notes, review = repair_prices(df)
        if notes:
            repairs[s] = {"notes": notes, "review": review}
    return out, repairs


def data_checks(data, missing, repairs):
    rows = []
    for s, _ in UNIVERSE:
        if s in missing or s not in data:
            rows.append({"symbol": s, "first_date": "-", "last_date": "-", "rows": 0, "repairs_made": "",
                         "big_jumps_left": 0, "jump_dates": "", "flag": "NO DATA"})
            continue
        c = data[s]["Close"]
        r = c / c.shift(1)
        jumps = r[(r > 1 + JUMP_FLAG) | (r < 1 / (1 + JUMP_FLAG))]
        rep = repairs.get(s)
        if len(jumps):
            flag = "CHECK — big one-day moves remain"
        elif rep and rep["review"]:
            flag = "CHECK — earlier history rescaled, verify"
        elif rep:
            flag = "REPAIRED — wrong prices corrected"
        else:
            flag = "ok"
        rows.append({"symbol": s, "first_date": c.index[0].date().isoformat(),
                     "last_date": c.index[-1].date().isoformat(), "rows": len(c),
                     "repairs_made": " | ".join(rep["notes"]) if rep else "",
                     "big_jumps_left": len(jumps),
                     "jump_dates": ", ".join(d.date().isoformat() for d in jumps.index[:6]), "flag": flag})
    return pd.DataFrame(rows)


# =============================================================================
# RANKS (point-in-time, same arithmetic as the live tracker)
# =============================================================================
def ranks_on(series_np, d):
    rows = []
    for s, (idx, close) in series_np.items():
        n = int(idx.searchsorted(d, side="right"))
        if n < MIN_HISTORY:
            continue
        if (d - idx[n - 1]).days > STALE_DAYS:
            continue
        cur = close[n - 1]
        rets = []
        for lb in LB:
            past = close[n - lb]
            if past <= 0:
                rets = None
                break
            rets.append((cur / past - 1.0) * 100.0)
        if rets is None:
            continue
        dma = float(close[n - DMA_PERIOD:n].mean())
        rows.append((s, sum(w * v for w, v in zip(WEIGHTS, rets)), rets[1], cur, dma))
    rows.sort(key=lambda r: -r[1])
    out = {}
    for i, (s, score, r3, cur, dma) in enumerate(rows, 1):
        out[s] = {"rank": i, "score": score, "ret3": r3, "price": cur, "dma": dma,
                  "elig": bool(cur > dma and r3 > 0)}
    return out


# =============================================================================
# BACKTEST ENGINE
# =============================================================================
def run_backtest(data, regime, bench):
    cp = COST_PCT / 100.0
    cal = regime.index
    first_ok = max(data[NG_NIFTY].index[0], data[NG_GOLD].index[0], cal[REGIME_DMA - 1])
    start = max(pd.Timestamp(START_DATE), first_ok)
    end = pd.Timestamp(END_DATE) if END_DATE else cal[-1]
    days = cal[(cal >= start) & (cal <= end)]
    if len(days) < 30:
        raise SystemExit("Not enough trading days in the chosen period.")

    syms = [s for s, _ in UNIVERSE if s in data]
    C = pd.DataFrame({s: data[s]["Close"] for s in syms}).reindex(cal).ffill().reindex(days)
    O = pd.DataFrame({s: data[s]["Open"] for s in syms}).reindex(days)
    H = pd.DataFrame({s: data[s]["High"] for s in syms}).reindex(days)
    L = pd.DataFrame({s: data[s]["Low"] for s in syms}).reindex(days)
    Cn = {s: C[s].values for s in syms}; On = {s: O[s].values for s in syms}
    Hn = {s: H[s].values for s in syms}; Ln = {s: L[s].values for s in syms}
    series_np = {s: (data[s].index, data[s]["Close"].values) for s in syms}
    reg_idx, reg_val = regime.index, regime.values

    ng_cap, mom_cap, reserve = CAPITAL * NG_SHARE, CAPITAL * MOM_SHARE, CAPITAL * CASH_SHARE

    # ---- Part A state ----
    ng = {"n": 0, "g": 0, "cash": ng_cap, "ref": None, "basis_n": 0.0, "basis_g": 0.0,
          "realised": 0.0, "costs": 0.0, "yield": 0.0}
    ng_log = []
    # ---- Part B state ----
    mom = {"cash": mom_cap, "yield": 0.0}
    hold, trades, monthly = {}, [], []
    regime_on_last = None

    rec = {k: [] for k in ["ng_value", "mom_value", "reserve", "total", "ratio", "gold_w",
                           "gold_w_actual", "n_hold", "regime", "rankable"]}
    rankable_last = 0
    prev_d = None

    for i, d in enumerate(days):
        # ---------- idle-cash yield ----------
        if prev_d is not None and CASH_YIELD_PA > 0:
            f = (1 + CASH_YIELD_PA) ** ((d - prev_d).days / 365.0) - 1
            ng["yield"] += ng["cash"] * f; ng["cash"] *= 1 + f
            mom["yield"] += mom["cash"] * f; mom["cash"] *= 1 + f
            reserve *= 1 + f

        # ================= Part A: Nifty:Gold ratio =================
        pn, pg = Cn[NG_NIFTY][i], Cn[NG_GOLD][i]
        ratio = pn / pg
        w = min(1.0, max(0.0, (ratio - NG_MIN) / (NG_MAX - NG_MIN)))
        if ng["ref"] is None or abs(ratio - ng["ref"]) >= NG_STEP:
            value = ng["n"] * pn + ng["g"] * pg + ng["cash"]
            tn, tg = int(value * (1 - w) / pn), int(value * w / pg)
            before = (ng["n"], ng["g"])
            cost_evt, traded = 0.0, 0.0
            for key, bkey, p, tgt in (("n", "basis_n", pn, tn), ("g", "basis_g", pg, tg)):   # sells first
                q = ng[key] - tgt
                if q > 0:
                    avg = ng[bkey] / ng[key]
                    c = q * p * cp + DP_CHARGE
                    ng["cash"] += q * p - c
                    ng["realised"] += q * (p - avg)
                    ng[bkey] -= q * avg
                    ng[key] -= q
                    cost_evt += c; traded += q * p
            for key, bkey, p, tgt in (("n", "basis_n", pn, tn), ("g", "basis_g", pg, tg)):   # then buys
                q = tgt - ng[key]
                if q > 0:
                    q = min(q, int(ng["cash"] / (p * (1 + cp))))
                    if q > 0:
                        c = q * p * cp
                        ng["cash"] -= q * p + c
                        ng[bkey] += q * p
                        ng[key] += q
                        cost_evt += c; traded += q * p
            ng["costs"] += cost_evt
            ng_log.append({"date": d.date().isoformat(),
                           "event": "Initial purchase" if ng["ref"] is None else "Reshuffle",
                           "ratio": ratio, "previous_ref_ratio": ng["ref"], "gold_weight_target_pct": w * 100,
                           "niftybees_price": pn, "goldbees_price": pg,
                           "niftybees_units_before": before[0], "niftybees_units_after": ng["n"],
                           "goldbees_units_before": before[1], "goldbees_units_after": ng["g"],
                           "value_traded": traded, "costs": cost_evt, "part_value_after":
                           ng["n"] * pn + ng["g"] * pg + ng["cash"]})
            ng["ref"] = ratio
        ng_value = ng["n"] * pn + ng["g"] * pg + ng["cash"]

        # ================= Part B: ETF momentum =================
        sold_today = set()

        def sell(s, price, reason):
            h = hold.pop(s)
            gross = h["qty"] * price
            c = gross * cp + DP_CHARGE
            mom["cash"] += gross - c
            net = h["qty"] * (price - h["entry"]) - h["cost_in"] - c
            trades.append({"symbol": s, "category": CATEGORY.get(s, ""), "entry_date": h["date"].date().isoformat(),
                           "entry_price": h["entry"], "qty": h["qty"], "exit_date": d.date().isoformat(),
                           "exit_price": price, "reason": reason, "rank_at_entry": h["rank"],
                           "days_held": (d - h["date"]).days,
                           "gross_pnl": h["qty"] * (price - h["entry"]), "costs": h["cost_in"] + c,
                           "net_pnl": net, "net_pnl_pct": net / (h["qty"] * h["entry"]) * 100})
            sold_today.add(s)

        # -- daily stops (the live strategy checks these on every tick) --
        for s in list(hold):
            h = hold[s]
            if h["date"] == d:
                continue
            sl_level = h["entry"] * (1 - SL_PCT)
            if STOP_BASIS == "low":
                lo, op, hi = Ln[s][i], On[s][i], Hn[s][i]
                if np.isnan(lo):
                    continue
                tr_level = h["peak"] * (1 - TRAIL_PCT)
                level = max(sl_level, tr_level)
                if lo <= level:
                    fill = level if (np.isnan(op) or op >= level) else op     # gap-down fills at the open
                    sell(s, float(fill), "Hard stop" if sl_level >= tr_level else "Trailing stop")
                elif not np.isnan(hi):
                    h["peak"] = max(h["peak"], hi)
            else:
                p = Cn[s][i]
                h["peak"] = max(h["peak"], p)
                if p <= sl_level:
                    sell(s, float(p), "Hard stop")
                elif p <= h["peak"] * (1 - TRAIL_PCT):
                    sell(s, float(p), "Trailing stop")

        # -- monthly rebalance: first trading day of a month (and the first backtest day) --
        if prev_d is None or d.month != prev_d.month:
            rk = ranks_on(series_np, d)
            rankable_last = len(rk)
            n = int(reg_idx.searchsorted(d, side="right"))
            regime_on = bool(reg_val[n - 1] > reg_val[n - REGIME_DMA:n].mean())
            regime_on_last = regime_on
            sells, buys = [f"{s} (stop)" for s in sold_today], []
            if not regime_on:
                for s in list(hold):
                    sell(s, float(Cn[s][i]), "Regime off"); sells.append(f"{s} (regime)")
            else:
                for s in list(hold):
                    r = rk.get(s)
                    if r is not None and r["rank"] > HOLD_RANK:
                        sell(s, float(Cn[s][i]), f"Rotation (rank {r['rank']})"); sells.append(f"{s} (rank {r['rank']})")
                value = mom["cash"] + sum(h["qty"] * Cn[s][i] for s, h in hold.items())
                slot = value / N_HOLD
                cands = sorted((r["rank"], s) for s, r in rk.items()
                               if r["rank"] <= N_HOLD and r["elig"] and s not in hold and s not in sold_today)
                for rank, s in cands:
                    if len(hold) >= N_HOLD:
                        break
                    p = float(Cn[s][i])
                    qty = int(min(slot, mom["cash"]) / (p * (1 + cp)))
                    if qty <= 0:
                        continue
                    c = qty * p * cp
                    mom["cash"] -= qty * p + c
                    hold[s] = {"qty": qty, "entry": p, "date": d, "peak": p, "cost_in": c, "rank": rank}
                    buys.append(f"{s} (rank {rank})")
            top6 = sorted((r["rank"], s, r["elig"]) for s, r in rk.items() if r["rank"] <= N_HOLD)
            monthly.append({"date": d.date().isoformat(), "regime": "ON" if regime_on else "OFF",
                            "etfs_rankable": len(rk),
                            "top6 (* = passes 3M>0 and 200-DMA)": ", ".join(f"{s}{'*' if e else ''}" for _, s, e in top6),
                            "sold": ", ".join(sells), "bought": ", ".join(buys),
                            "holdings_after": ", ".join(sorted(hold)), "n_holdings": len(hold),
                            "cash_after": mom["cash"],
                            "part_b_value": mom["cash"] + sum(h["qty"] * Cn[s][i] for s, h in hold.items())})

        mom_value = mom["cash"] + sum(h["qty"] * Cn[s][i] for s, h in hold.items())
        rec["ng_value"].append(ng_value); rec["mom_value"].append(mom_value); rec["reserve"].append(reserve)
        rec["total"].append(ng_value + mom_value + reserve)
        rec["ratio"].append(ratio); rec["gold_w"].append(w * 100)
        rec["gold_w_actual"].append(ng["g"] * pg / ng_value * 100 if ng_value else 0)
        rec["n_hold"].append(len(hold)); rec["regime"].append(1 if regime_on_last else 0)
        rec["rankable"].append(rankable_last)
        prev_d = d

    eq = pd.DataFrame(rec, index=days)
    last_i = len(days) - 1
    pn, pg = Cn[NG_NIFTY][last_i], Cn[NG_GOLD][last_i]

    # benchmarks, rebased to the starting capital
    b = bench.reindex(cal).ffill().reindex(days) if bench is not None else None
    eq["bench"] = (b / b.iloc[0] * CAPITAL) if b is not None and b.notna().all() else np.nan
    eq["niftybees_bh"] = C[NG_NIFTY] / C[NG_NIFTY].iloc[0] * CAPITAL

    open_pos = pd.DataFrame([{
        "symbol": s, "category": CATEGORY.get(s, ""), "entry_date": h["date"].date().isoformat(),
        "entry_price": h["entry"], "qty": h["qty"], "last_price": float(Cn[s][last_i]),
        "peak": h["peak"], "value": h["qty"] * float(Cn[s][last_i]),
        "unrealised_net": h["qty"] * (float(Cn[s][last_i]) - h["entry"]) - h["cost_in"],
        "pnl_pct": (float(Cn[s][last_i]) / h["entry"] - 1) * 100} for s, h in hold.items()])

    # ---- reconciliation: rebuild each part's end value from its trade records ----
    tr = pd.DataFrame(trades)
    closed_net = float(tr["net_pnl"].sum()) if len(tr) else 0.0
    open_net = float(open_pos["unrealised_net"].sum()) if len(open_pos) else 0.0
    recon = [
        ("Part B — start capital", mom_cap), ("Part B — net P&L of closed trades", closed_net),
        ("Part B — unrealised P&L of open positions (net of entry cost)", open_net),
        ("Part B — interest on idle cash", mom["yield"]),
        ("Part B — end value rebuilt from trades", mom_cap + closed_net + open_net + mom["yield"]),
        ("Part B — end value from the daily engine", float(eq["mom_value"].iloc[-1])),
        ("Part A — start capital", ng_cap), ("Part A — realised P&L on units sold", ng["realised"]),
        ("Part A — unrealised P&L on units held",
         ng["n"] * pn + ng["g"] * pg - ng["basis_n"] - ng["basis_g"]),
        ("Part A — costs paid", -ng["costs"]), ("Part A — interest on idle cash", ng["yield"]),
        ("Part A — end value rebuilt from trades",
         ng_cap + ng["realised"] + (ng["n"] * pn + ng["g"] * pg - ng["basis_n"] - ng["basis_g"])
         - ng["costs"] + ng["yield"]),
        ("Part A — end value from the daily engine", float(eq["ng_value"].iloc[-1])),
        ("Idle cash reserve", float(eq["reserve"].iloc[-1])),
        ("TOTAL — end value from the daily engine", float(eq["total"].iloc[-1])),
    ]
    state = {"ng": ng, "ng_cap": ng_cap, "mom_cap": mom_cap, "pn": pn, "pg": pg,
             "mom_cash": mom["cash"], "mom_costs": float(tr["costs"].sum()) if len(tr) else 0.0}
    return eq, tr, open_pos, pd.DataFrame(ng_log), pd.DataFrame(monthly), recon, state


# =============================================================================
# STATISTICS
# =============================================================================
def stats(s):
    s = s.dropna()
    if len(s) < 2:
        return None
    yrs = max((s.index[-1] - s.index[0]).days / 365.25, 1e-9)
    ret = s.pct_change().dropna()
    cagr = (s.iloc[-1] / s.iloc[0]) ** (1 / yrs) - 1
    vol = float(ret.std() * math.sqrt(252))
    dd = s / s.cummax() - 1
    m = period_returns(s, "M")
    return {"start_value": float(s.iloc[0]), "end_value": float(s.iloc[-1]),
            "total_return_pct": (s.iloc[-1] / s.iloc[0] - 1) * 100, "cagr_pct": cagr * 100,
            "volatility_pct": vol * 100, "sharpe": ((cagr - RF_PA) / vol if vol > 0 else float("nan")),
            "max_drawdown_pct": float(dd.min()) * 100, "max_drawdown_date": dd.idxmin().date().isoformat(),
            "calmar": (cagr / abs(dd.min()) if dd.min() < 0 else float("nan")),
            "best_month_pct": float(m.max()) * 100, "worst_month_pct": float(m.min()) * 100,
            "positive_months_pct": float((m > 0).mean()) * 100, "years": yrs}


def period_returns(s, freq):
    """Returns per calendar month ('M') or year ('Y'), first period measured from the start value."""
    s = s.dropna()
    last = s.groupby(s.index.to_period(freq)).last()
    prev = last.shift(1)
    prev.iloc[0] = s.iloc[0]
    return last / prev - 1


# =============================================================================
# EXCEL  (written with openpyxl directly)
# =============================================================================
def build_excel(eq, tr, open_pos, ng_log, monthly, recon, checks, summary_rows, yearly, mtable, assumptions, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    hdr = Font(bold=True, color="FFFFFF"); hdrfill = PatternFill("solid", fgColor="1F3B57")
    green = PatternFill("solid", fgColor="C6EFCE"); red = PatternFill("solid", fgColor="FFC7CE")
    amber = PatternFill("solid", fgColor="FFEB9C")
    wb = Workbook(); wb.remove(wb.active)

    def clean(v):
        if v is None:
            return None
        if isinstance(v, (np.floating, float)):
            return None if (math.isnan(v) or math.isinf(v)) else round(float(v), 4)
        if isinstance(v, np.integer):
            return int(v)
        if isinstance(v, (pd.Timestamp, datetime)):
            return v.strftime("%Y-%m-%d")
        return v

    def sheet(name, df):
        ws = wb.create_sheet(name)
        ws.append([str(c) for c in df.columns])
        for row in df.itertuples(index=False):
            ws.append([clean(v) for v in row])
        for cell in ws[1]:
            cell.font = hdr; cell.fill = hdrfill; cell.alignment = Alignment(horizontal="center")
        for col in ws.columns:
            wdt = max((len(str(c.value)) for c in col[:200] if c.value is not None), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(max(wdt + 2, 10), 60)
        ws.freeze_panes = "A2"
        return ws

    sheet("Summary", pd.DataFrame(summary_rows))
    sheet("Yearly_Returns", yearly)
    sheet("Monthly_Returns", mtable)
    ws = sheet("Momentum_Trades", tr if len(tr) else pd.DataFrame(columns=["symbol"]))
    if len(tr):
        col = list(tr.columns).index("net_pnl") + 1
        for r in range(2, ws.max_row + 1):
            v = ws.cell(r, col).value
            fill = green if (v or 0) > 0 else red
            for c in range(1, len(tr.columns) + 1):
                ws.cell(r, c).fill = fill
    sheet("Momentum_Open", open_pos if len(open_pos) else pd.DataFrame(columns=["symbol"]))
    sheet("Momentum_Monthly", monthly)
    sheet("NiftyGold_Reshuffles", ng_log)
    d = eq.copy(); d.insert(0, "date", [x.date().isoformat() for x in d.index])
    sheet("Daily_Equity", d)

    # Reconciliation, with live formulas for the differences
    ws = wb.create_sheet("Reconciliation")
    ws.append(["Item", "Rs"])
    for cell in ws[1]:
        cell.font = hdr; cell.fill = hdrfill
    for k, v in recon:
        ws.append([k, round(float(v), 2)])
    pos = {k: i + 2 for i, (k, _) in enumerate(recon)}
    r0 = ws.max_row + 2
    ws.cell(r0, 1, "Part B difference (rebuilt − engine), should be 0")
    ws.cell(r0, 2, f"=ROUND(B{pos['Part B — end value rebuilt from trades']}-B{pos['Part B — end value from the daily engine']},2)")
    ws.cell(r0 + 1, 1, "Part A difference (rebuilt − engine), should be 0")
    ws.cell(r0 + 1, 2, f"=ROUND(B{pos['Part A — end value rebuilt from trades']}-B{pos['Part A — end value from the daily engine']},2)")
    ws.cell(r0 + 2, 1, "Total = Part A + Part B + reserve, difference should be 0")
    ws.cell(r0 + 2, 2, f"=ROUND(B{pos['Part A — end value from the daily engine']}+B{pos['Part B — end value from the daily engine']}"
                       f"+B{pos['Idle cash reserve']}-B{pos['TOTAL — end value from the daily engine']},2)")
    for r in (r0, r0 + 1, r0 + 2):
        ws.cell(r, 1).font = Font(bold=True); ws.cell(r, 2).fill = amber
    ws.column_dimensions["A"].width = 66; ws.column_dimensions["B"].width = 18

    ws = sheet("Data_Checks", checks)
    fcol = list(checks.columns).index("flag") + 1
    for r in range(2, ws.max_row + 1):
        v = str(ws.cell(r, fcol).value)
        if v != "ok":
            for c in range(1, len(checks.columns) + 1):
                ws.cell(r, c).fill = amber if v.startswith("REPAIRED") else red
    sheet("Assumptions", pd.DataFrame({"Assumptions and limits — read before trusting the numbers": assumptions}))
    wb.save(path)


# =============================================================================
# HTML
# =============================================================================
DARK = dict(paper_bgcolor="#0D1B2A", plot_bgcolor="#0D1B2A",
            font=dict(color="#E9EEF7"), margin=dict(l=50, r=20, t=44, b=40),
            legend=dict(orientation="h", y=-0.16), xaxis=dict(gridcolor="#24384f"),
            yaxis=dict(gridcolor="#24384f"))
_first = [True]


def _div(fig):
    if _first[0]:
        _first[0] = False
        return pyo.plot(fig, include_plotlyjs=True, output_type="div")
    return pyo.plot(fig, include_plotlyjs=False, output_type="div")


def build_html(eq, tr, open_pos, ng_log, st, yearly, mret, checks, assumptions, meta, path):
    acc, pos, neg, mut, blu = "#F5A524", "#37D3A6", "#E5595E", "#8B98AD", "#5AA9E6"
    has_bench = eq["bench"].notna().all()

    def rebase(s):
        return s / s.iloc[0] * 100

    charts = []
    f = go.Figure()
    f.add_trace(go.Scatter(x=eq.index, y=rebase(eq["total"]), name="Combined strategy", line=dict(color=acc, width=2.4)))
    f.add_trace(go.Scatter(x=eq.index, y=rebase(eq["ng_value"]), name="Part A — Nifty:Gold", line=dict(color=pos, width=1.3)))
    f.add_trace(go.Scatter(x=eq.index, y=rebase(eq["mom_value"]), name="Part B — ETF momentum", line=dict(color=blu, width=1.3)))
    if has_bench:
        f.add_trace(go.Scatter(x=eq.index, y=rebase(eq["bench"]), name=BENCH_LABEL, line=dict(color=mut, width=1.3, dash="dot")))
    f.update_layout(title="Growth of 100 (log scale)", yaxis_type="log", **DARK)
    charts.append(_div(f))

    f = go.Figure()
    f.add_trace(go.Scatter(x=eq.index, y=(eq["total"] / eq["total"].cummax() - 1) * 100, name="Combined strategy",
                           fill="tozeroy", line=dict(color=neg, width=1.2)))
    if has_bench:
        f.add_trace(go.Scatter(x=eq.index, y=(eq["bench"] / eq["bench"].cummax() - 1) * 100, name=BENCH_LABEL,
                               line=dict(color=mut, width=1, dash="dot")))
    f.update_layout(title="Drawdown from the previous peak (%)", **DARK)
    charts.append(_div(f))

    f = go.Figure()
    f.add_trace(go.Bar(x=yearly["year"], y=yearly["combined_pct"], name="Combined strategy", marker_color=acc))
    if has_bench:
        f.add_trace(go.Bar(x=yearly["year"], y=yearly["benchmark_pct"], name=BENCH_LABEL, marker_color=mut))
    f.update_layout(title="Calendar-year returns (%) — first and last years are part-years", barmode="group", **DARK)
    charts.append(_div(f))

    mm = (mret * 100).round(1)
    piv = pd.DataFrame({"y": mm.index.year, "m": mm.index.month, "v": mm.values}).pivot(index="y", columns="m", values="v")
    piv = piv.reindex(columns=range(1, 13))
    f = go.Figure(go.Heatmap(z=piv.values, x=["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"],
                             y=[str(y) for y in piv.index], colorscale=[[0, neg], [0.5, "#12263A"], [1, pos]], zmid=0,
                             text=[["" if pd.isna(v) else f"{v:.1f}" for v in row] for row in piv.values],
                             texttemplate="%{text}", textfont=dict(size=10), colorbar=dict(title="%")))
    f.update_layout(title="Monthly returns of the combined strategy (%)",
                    **{**DARK, "xaxis": dict(showgrid=False, zeroline=False),
                       "yaxis": dict(autorange="reversed", showgrid=False, zeroline=False)})
    charts.append(_div(f))

    f = go.Figure()
    f.add_trace(go.Scatter(x=eq.index, y=eq["ratio"], name="NIFTYBEES / GOLDBEES ratio", line=dict(color=acc, width=1.6)))
    f.add_trace(go.Scatter(x=eq.index, y=eq["gold_w_actual"], name="Gold weight actually held (%)", yaxis="y2",
                           line=dict(color=pos, width=1.3)))
    if len(ng_log):
        f.add_trace(go.Scatter(x=pd.to_datetime(ng_log["date"]), y=ng_log["ratio"], mode="markers", name="Purchase / reshuffle",
                               marker=dict(color="#FFFFFF", size=9, symbol="diamond")))
    f.update_layout(title="Part A — ratio (left), reshuffle points and gold weight % (right)",
                    yaxis2=dict(overlaying="y", side="right", range=[0, 100], showgrid=False, tickmode="array",
                                tickvals=[0, 20, 40, 60, 80, 100], ticksuffix="%"), **DARK)
    charts.append(_div(f))

    f = go.Figure()
    f.add_trace(go.Scatter(x=eq.index, y=eq["n_hold"], name="ETFs held (max 6)", line=dict(color=blu, width=1.3, shape="hv")))
    f.add_trace(go.Scatter(x=eq.index, y=eq["rankable"], name="ETFs with enough history to rank", yaxis="y2",
                           line=dict(color=mut, width=1.2, dash="dot", shape="hv")))
    f.update_layout(title="Part B — ETFs held (left), and how many of the 50 ETFs had enough history to rank (right)",
                    yaxis=dict(range=[0, 6.5], gridcolor="#24384f"),
                    yaxis2=dict(overlaying="y", side="right", range=[0, 52], showgrid=False, tickmode="array",
                                tickvals=[0, 10, 20, 30, 40, 50]),
                    **{k: v for k, v in DARK.items() if k != "yaxis"})
    charts.append(_div(f))

    c = st["Combined"]
    kpis = [("Final value ₹", f"{c['end_value']:,.0f}"), ("CAGR", f"{c['cagr_pct']:.1f}%"),
            ("Max drawdown", f"{c['max_drawdown_pct']:.1f}%"), ("Sharpe", f"{c['sharpe']:.2f}"),
            ("Momentum trades", f"{len(tr)}"), ("Gold reshuffles", f"{max(len(ng_log) - 1, 0)}")]
    kpi_html = "".join(f"<div class='kpi'><div class='kv'>{v}</div><div class='kl'>{k}</div></div>" for k, v in kpis)

    rows_def = [("Start value ₹", "start_value", "{:,.0f}"), ("End value ₹", "end_value", "{:,.0f}"),
                ("Total return", "total_return_pct", "{:+.1f}%"), ("CAGR", "cagr_pct", "{:+.1f}%"),
                ("Volatility (yearly)", "volatility_pct", "{:.1f}%"), (f"Sharpe (risk-free {RF_PA*100:.0f}%)", "sharpe", "{:.2f}"),
                ("Max drawdown", "max_drawdown_pct", "{:.1f}%"), ("Max drawdown date", "max_drawdown_date", "{}"),
                ("CAGR ÷ max drawdown", "calmar", "{:.2f}"), ("Best month", "best_month_pct", "{:+.1f}%"),
                ("Worst month", "worst_month_pct", "{:+.1f}%"), ("Positive months", "positive_months_pct", "{:.0f}%")]
    cols = [k for k in st if st[k] is not None]
    stat_tbl = ("<div class='tw'><table><thead><tr><th style='text-align:left'>Measure</th>"
                + "".join(f"<th>{k}</th>" for k in cols) + "</tr></thead><tbody>"
                + "".join("<tr><td class='sym'>" + lab + "</td>" + "".join(
                    f"<td class='num'>{fmt.format(st[k][key])}</td>" for k in cols) + "</tr>"
                    for lab, key, fmt in rows_def) + "</tbody></table></div>")

    def tbl(df, cols, tid, pct_cols=(), money_cols=()):
        if not len(df):
            return "<div class='empty'>None.</div>"
        head = "".join(f"<th>{c.replace('_', ' ')}</th>" for c in cols)
        body = ""
        for r in df[cols].itertuples(index=False):
            tds = ""
            for cn, v in zip(cols, r):
                if cn in pct_cols:
                    tds += f"<td class='num {'pos' if v >= 0 else 'neg'}'>{v:+.1f}%</td>"
                elif cn in money_cols:
                    tds += f"<td class='num {'pos' if v >= 0 else 'neg'}'>{v:,.0f}</td>"
                elif isinstance(v, (float, np.floating)):
                    tds += f"<td class='num'>{v:,.2f}</td>"
                elif isinstance(v, (int, np.integer)):
                    tds += f"<td class='num'>{v:,}</td>"
                else:
                    tds += f"<td class='l'>{'' if v is None else v}</td>"
            body += f"<tr>{tds}</tr>"
        return f"<div class='tw' style='max-height:460px;overflow-y:auto'><table id='{tid}' class='sortable'><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table></div>"

    trade_tbl = tbl(tr.iloc[::-1] if len(tr) else tr,
                    ["symbol", "entry_date", "entry_price", "qty", "exit_date", "exit_price", "reason",
                     "days_held", "net_pnl", "net_pnl_pct"], "t1", pct_cols=("net_pnl_pct",), money_cols=("net_pnl",))
    open_tbl = tbl(open_pos, ["symbol", "entry_date", "entry_price", "qty", "last_price", "value", "pnl_pct"], "t2",
                   pct_cols=("pnl_pct",))
    ng_tbl = tbl(ng_log, ["date", "event", "ratio", "gold_weight_target_pct", "niftybees_units_after",
                          "goldbees_units_after", "value_traded", "costs", "part_value_after"], "t3")

    if len(tr):
        wins = tr[tr["net_pnl"] > 0]
        by_reason = tr.assign(kind=tr["reason"].str.replace(r" \(rank \d+\)", "", regex=True)).groupby("kind").agg(
            trades=("net_pnl", "size"), net_pnl=("net_pnl", "sum"), avg_pct=("net_pnl_pct", "mean")).reset_index()
        trade_line = (f"{len(tr)} closed trades · {len(wins) / len(tr) * 100:.0f}% profitable · average "
                      f"{tr['net_pnl_pct'].mean():+.1f}% · average holding {tr['days_held'].mean():.0f} days. By exit type: "
                      + " · ".join(f"{r.kind}: {r.trades} trades, ₹{r.net_pnl:,.0f}" for r in by_reason.itertuples()))
    else:
        trade_line = "No closed trades in this period."

    bad = checks[checks["flag"].str.startswith("CHECK") | (checks["flag"] == "NO DATA")]
    fixed = checks[checks["flag"].str.startswith("REPAIRED")]
    warn = ""
    if len(bad):
        warn += ("<div class='regime off'><b>Data warnings on " + str(len(bad)) + " ETF(s)</b> · "
                 + "; ".join(f"{r.symbol}: {r.flag}" + (f" ({r.jump_dates})" if r.jump_dates else "")
                             + (f" [{r.repairs_made}]" if r.repairs_made else "") for r in bad.itertuples())
                 + "<span class='sm'>A flagged ETF can distort ranks, stops and returns. See the Data_Checks sheet.</span></div>")
    if len(fixed):
        warn += ("<div class='regime fix'><b>Wrong Yahoo prices repaired on " + str(len(fixed)) + " ETF(s)</b> · "
                 + "; ".join(f"{r.symbol}: {r.repairs_made}" for r in fixed.itertuples())
                 + "<span class='sm'>These were split errors in the data, not real price moves. Details in the Data_Checks sheet.</span></div>")
    if meta.get("mock"):
        warn = "<div class='regime off'><b>MOCK DATA</b> · synthetic prices, for testing the script only. These are not real results.</div>" + warn

    css = """
    :root{--bg:#0D1B2A;--panel:#12263A;--line:#24384f;--fg:#E9EEF7;--fg2:#8B98AD;--acc:#F5A524;--pos:#37D3A6;--neg:#E5595E}
    *{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--fg);font-family:Inter,system-ui,Arial,sans-serif}
    .wrap{max-width:1180px;margin:0 auto;padding:26px 20px 70px}
    h1{font-size:20px;margin:0 0 2px}.sub{color:var(--fg2);font-size:13px;margin-bottom:16px;line-height:1.5}
    h2{font-size:13px;text-transform:uppercase;letter-spacing:.06em;color:var(--fg2);margin:26px 0 10px}
    .kpis{display:grid;grid-template-columns:repeat(6,1fr);gap:12px}
    @media(max-width:800px){.kpis{grid-template-columns:repeat(2,1fr)}}
    .kpi{background:var(--panel);border:1px solid var(--line);border-radius:12px;padding:14px}
    .kpi .kv{font-size:20px;font-weight:600}.kpi .kl{color:var(--fg2);font-size:11px;text-transform:uppercase;letter-spacing:.05em;margin-top:4px}
    .empty{background:var(--panel);border:1px solid var(--line);border-radius:12px;padding:16px;color:var(--fg2)}
    .tw{overflow-x:auto;border:1px solid var(--line);border-radius:12px;margin-top:6px}
    table{border-collapse:collapse;width:100%;background:var(--panel)}
    th{position:sticky;top:0;background:#183049;color:var(--fg2);font-size:10.5px;text-transform:uppercase;letter-spacing:.04em;padding:9px 10px;text-align:right;border-bottom:1px solid var(--line);white-space:nowrap}
    table.sortable th{cursor:pointer}
    td{padding:8px 10px;border-bottom:1px solid rgba(36,56,79,.6);text-align:right;font-size:12.5px;white-space:nowrap}
    td.sym,td.l{text-align:left}td.sym{font-weight:600}
    .num{font-variant-numeric:tabular-nums}.pos{color:var(--pos)}.neg{color:var(--neg)}
    .rules{background:var(--panel);border:1px solid var(--line);border-radius:12px;padding:14px 16px;color:var(--fg2);font-size:12.5px;line-height:1.7}
    .rules li{margin-bottom:4px}.rules b{color:var(--fg)}
    .line{color:var(--fg2);font-size:12.5px;margin:6px 0 8px;line-height:1.6}
    .chart{background:var(--panel);border:1px solid var(--line);border-radius:12px;margin:12px 0;padding:6px}
    .regime{border-radius:12px;padding:13px 16px;margin:14px 0 6px;font-size:13px;border:1px solid var(--line);background:var(--panel);line-height:1.5}
    .regime .sm{display:block;font-size:11px;color:var(--fg2);margin-top:5px}
    .regime.fix{border-color:rgba(245,165,36,.5);background:linear-gradient(160deg,rgba(245,165,36,.12),rgba(245,165,36,.03))}.regime.fix b{color:var(--acc)}
    .regime.off{border-color:rgba(229,89,94,.55);background:linear-gradient(160deg,rgba(229,89,94,.16),rgba(229,89,94,.03))}.regime.off b{color:var(--neg)}
    """
    sort_js = """
    <script>
    document.querySelectorAll('table.sortable').forEach(function(t){t.querySelectorAll('th').forEach(function(th,i){
      th.addEventListener('click',function(){var tb=t.querySelector('tbody');var rows=[].slice.call(tb.querySelectorAll('tr'));
      var asc=th._asc=!th._asc;rows.sort(function(a,b){var x=a.children[i].innerText,y=b.children[i].innerText;
        var nx=parseFloat(x.replace(/[^0-9.\\-]/g,'')),ny=parseFloat(y.replace(/[^0-9.\\-]/g,''));
        if(!isNaN(nx)&&!isNaN(ny)&&!/^\\d{4}-\\d{2}-\\d{2}$/.test(x))return asc?nx-ny:ny-nx;
        return asc?x.localeCompare(y):y.localeCompare(x);});rows.forEach(function(r){tb.appendChild(r);});});});});
    </script>"""

    header = (f"<h1>Nifty:Gold Ratio + ETF Momentum — Combined Backtest</h1>"
              f"<div class='sub'>{meta['start']} to {meta['end']} · {c['years']:.1f} years · start capital ₹{CAPITAL:,.0f} "
              f"({NG_SHARE*100:.0f}% Nifty:Gold, {MOM_SHARE*100:.0f}% ETF momentum, {CASH_SHARE*100:.0f}% cash) · "
              f"stops checked on the daily <b>{STOP_BASIS}</b> · costs {COST_PCT:.2f}% each side + ₹{DP_CHARGE:.2f} per sell · "
              f"generated {meta['generated']}</div>")
    html = ("<!DOCTYPE html><html lang='en'><head><meta charset='utf-8'>"
            "<meta name='viewport' content='width=device-width, initial-scale=1'>"
            "<title>Nifty:Gold + ETF Momentum Backtest</title>"
            "<link href='https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap' rel='stylesheet'>"
            f"<style>{css}</style></head><body><div class='wrap'>{header}{warn}"
            f"<div class='kpis'>{kpi_html}</div>"
            f"<h2>Results</h2>{stat_tbl}"
            + "".join(f"<div class='chart'>{ch}</div>" for ch in charts) +
            f"<h2>Part B — closed momentum trades (click a heading to sort)</h2><div class='line'>{trade_line}</div>{trade_tbl}"
            f"<h2>Part B — positions open at the end</h2>{open_tbl}"
            f"<h2>Part A — purchase and reshuffles</h2>{ng_tbl}"
            "<h2>Assumptions and limits — read before trusting the numbers</h2><div class='rules'><ul>"
            + "".join(f"<li>{a}</li>" for a in assumptions) + "</ul></div>"
            f"</div>{sort_js}</body></html>")
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(html)


# =============================================================================
# MAIN
# =============================================================================
def main():
    print("Nifty:Gold Ratio + ETF Momentum — Combined Backtest")
    print("-" * 52)
    if STOP_BASIS not in ("close", "low"):
        raise SystemExit("BT_STOP_BASIS must be 'close' or 'low'.")
    print("[1] Loading prices    → ", end="")
    data, regime, bench, missing = load_data()
    for need in (NG_NIFTY, NG_GOLD):
        if need not in data:
            raise SystemExit(f"No price data for {need} — cannot run the backtest.")
    if regime is None or len(regime) < REGIME_DMA + 30:
        raise SystemExit(f"No usable {REGIME_LABEL} data — cannot run the backtest.")
    data, repairs = repair_all(data)
    checks = data_checks(data, missing, repairs)
    n_rep = int(checks["flag"].str.startswith("REPAIRED").sum())
    n_bad = int((checks["flag"].str.startswith("CHECK") | (checks["flag"] == "NO DATA")).sum())
    print(f"[2] Data checks       → {n_rep} ETF(s) had wrong prices repaired, {n_bad} need a look"
          + (" — see Data_Checks sheet" if (n_rep or n_bad) else ""))
    for sym, rep in repairs.items():
        for note in rep["notes"]:
            print(f"      {sym}: {note}")

    print("[3] Running backtest  → ", end="")
    eq, tr, open_pos, ng_log, monthly, recon, state = run_backtest(data, regime, bench)
    print(f"{eq.index[0].date()} to {eq.index[-1].date()}, {len(eq)} trading days, "
          f"{len(tr)} momentum trades, {max(len(ng_log) - 1, 0)} gold reshuffles")

    # Measure the strategy from the capital put in (before day-1 costs), not from the day-1 closing value.
    day0 = eq.index[0] - pd.Timedelta(days=1)

    def from_capital(col, cap):
        return pd.concat([pd.Series([cap], index=[day0]), eq[col]])

    s_tot = from_capital("total", CAPITAL)
    s_a = from_capital("ng_value", state["ng_cap"])
    s_b = from_capital("mom_value", state["mom_cap"])
    st = {"Combined": stats(s_tot), "Part A Nifty:Gold": stats(s_a),
          "Part B ETF momentum": stats(s_b), BENCH_LABEL: stats(eq["bench"]),
          "NIFTYBEES buy & hold": stats(eq["niftybees_bh"])}
    rd = dict(recon)
    diff_b = rd["Part B — end value rebuilt from trades"] - rd["Part B — end value from the daily engine"]
    diff_a = rd["Part A — end value rebuilt from trades"] - rd["Part A — end value from the daily engine"]
    print(f"[4] Reconciliation    → Part A diff ₹{diff_a:.2f}, Part B diff ₹{diff_b:.2f} (both should be 0)")

    def by_year(series):          # label by the year of each trading day (day 0 belongs to the first year)
        r = period_returns(series, "Y")
        return r[r.index >= eq.index[0].to_period("Y")]

    yr = {"combined_pct": by_year(s_tot), "part_a_pct": by_year(s_a),
          "part_b_pct": by_year(s_b), "niftybees_pct": by_year(eq["niftybees_bh"])}
    if eq["bench"].notna().all():
        yr["benchmark_pct"] = by_year(eq["bench"])
    yearly = (pd.DataFrame(yr) * 100).round(2)
    yearly.insert(0, "year", [str(p) for p in yearly.index])
    if "benchmark_pct" not in yearly:
        yearly["benchmark_pct"] = np.nan
    mret = period_returns(s_tot, "M")
    mret = mret[mret.index >= eq.index[0].to_period("M")]
    mret.index = mret.index.to_timestamp()
    mtable = pd.DataFrame({"year": mret.index.year, "month": mret.index.month, "v": (mret.values * 100).round(2)}) \
        .pivot(index="year", columns="month", values="v").reindex(columns=range(1, 13))
    mtable.columns = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    mtable.insert(0, "year", mtable.index)

    in_cash = float((eq["n_hold"] == 0).mean() * 100)
    first_rank = int(eq["rankable"].iloc[0]); last_rank = int(eq["rankable"].iloc[-1])
    assumptions = [
        f"<b>Survivorship.</b> The universe is today's list of 50 ETFs. Only {first_rank} of them had enough history to be "
        f"ranked at the start, rising to {last_rank} at the end, so the early years test a much narrower strategy than the one you will trade.",
        "<b>Prices.</b> Yahoo Finance daily data, adjusted for splits and dividends. Yahoo data for thinly traded ETFs can contain "
        "wrong prices, especially around unit splits. Any one-day move above "
        f"{JUMP_FLAG*100:.0f}% is treated as a data error and repaired; every repair is listed in the Data_Checks sheet. "
        "Smaller errors would not be caught.",
        f"<b>Stops.</b> Checked once a day on the daily {STOP_BASIS}. "
        + ("Live, Tradetron checks every tick, so live stops can trigger on intraday dips that a close-based test never sees. "
           "Re-run with stop basis 'low' to see the stricter version."
           if STOP_BASIS == "close" else
           "A stop is filled at the stop level, or at the day's open if the ETF gapped below it. The order of the day's high and low is unknown, "
           "so the trailing level uses the peak up to the previous day."),
        "<b>Fills.</b> Monthly ranks and trades both use the first trading day's closing price. Live, ranks are computed around 2:30 pm "
        "and orders go between 3:00 and 3:20 pm at market prices.",
        "<b>Nifty:Gold trigger.</b> The 1.0 ratio move is tested on daily closes. Live, it is tested through the day.",
        f"<b>Costs.</b> {COST_PCT:.2f}% of trade value on each buy and sell (charges and slippage together) plus ₹{DP_CHARGE:.2f} per sell. "
        "Real slippage on thinly traded ETFs can be higher. No taxes are deducted.",
        f"<b>Cash.</b> Idle cash earns {CASH_YIELD_PA*100:.1f}% a year. Part B was fully in cash on {in_cash:.0f}% of days.",
        "<b>Sizing.</b> Whole units only, at the starting capital you entered. Each part compounds on its own; "
        "there is no rebalancing between Part A and Part B.",
        "<b>Lookback.</b> Returns use the same arithmetic as the live tracker (1M = 21 rows of price history, and so on), "
        "so the backtest and the live signals rank the same way.",
        "<b>This is a historical simulation, not a forecast.</b> Past results do not predict future returns. Not investment advice.",
    ]
    plain = [a.replace("<b>", "").replace("</b>", "") for a in assumptions]

    summary_rows = []
    labels = [("Start value (Rs)", "start_value"), ("End value (Rs)", "end_value"), ("Total return %", "total_return_pct"),
              ("CAGR %", "cagr_pct"), ("Volatility % (yearly)", "volatility_pct"), (f"Sharpe (rf {RF_PA*100:.0f}%)", "sharpe"),
              ("Max drawdown %", "max_drawdown_pct"), ("Max drawdown date", "max_drawdown_date"),
              ("CAGR / max drawdown", "calmar"), ("Best month %", "best_month_pct"), ("Worst month %", "worst_month_pct"),
              ("Positive months %", "positive_months_pct")]
    for lab, key in labels:
        row = {"Measure": lab}
        for k, v in st.items():
            row[k] = (v[key] if v is not None else None)
            if isinstance(row[k], float):
                row[k] = round(row[k], 2)
        summary_rows.append(row)
    for lab, val in [("Period", f"{eq.index[0].date()} to {eq.index[-1].date()}"), ("Stop basis", STOP_BASIS),
                     ("Cost % each side", COST_PCT), ("DP charge per sell (Rs)", DP_CHARGE),
                     ("Momentum trades closed", len(tr)), ("Gold reshuffles", max(len(ng_log) - 1, 0)),
                     ("Total costs paid — Part A (Rs)", round(state["ng"]["costs"], 2)),
                     ("Total costs paid — Part B closed trades (Rs)", round(state["mom_costs"], 2)),
                     ("ETFs with wrong prices repaired", n_rep), ("ETFs needing a manual look", n_bad), ("Generated", datetime.now().strftime("%d %b %Y %H:%M"))]:
        summary_rows.append({"Measure": lab, "Combined": val})

    meta = {"start": eq.index[0].strftime("%d %b %Y"), "end": eq.index[-1].strftime("%d %b %Y"),
            "generated": datetime.now().strftime("%d %b %Y %H:%M"), "mock": os.getenv("BT_MOCK") == "1"}
    c = st["Combined"]
    print(f"    Combined: ₹{c['start_value']:,.0f} → ₹{c['end_value']:,.0f} | CAGR {c['cagr_pct']:.1f}% | "
          f"max drawdown {c['max_drawdown_pct']:.1f}% | Sharpe {c['sharpe']:.2f}")
    print("[5] Excel saved       → ", end="")
    build_excel(eq, tr, open_pos, ng_log, monthly, recon, checks, summary_rows, yearly, mtable, plain, EXCEL_PATH)
    print(EXCEL_PATH.name, "✓")
    print("[6] HTML saved        → ", end="")
    build_html(eq, tr, open_pos, ng_log, st, yearly, mret, checks, assumptions, meta, HTML_PATH)
    print(HTML_PATH.name, "✓")
    print("[7] Opening browser   → ", end="")
    try:
        if os.getenv("CI") or os.getenv("GITHUB_ACTIONS"):
            raise RuntimeError
        webbrowser.open(HTML_PATH.as_uri()); print("report launched ✓")
    except Exception:
        print("no browser (headless) — open the HTML from the output folder")
    print(f"\nSaved: {EXCEL_PATH}\nSaved: {HTML_PATH}")


if __name__ == "__main__":
    main()
